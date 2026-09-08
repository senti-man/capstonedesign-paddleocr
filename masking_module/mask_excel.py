import openpyxl
import os
from detector import detect_sensitive_items_fullmatch

# 엑셀 파일에서 텍스트를 추출하고 감지된 개인정보 항목을 반환

def extract_excel_sensitive_info(excel_path):
    '''
    개인정보 항목 탐지, 반환
    '''
    workbook = openpyxl.load_workbook(excel_path)
    extracted_text = []
    detected_info = {}

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]

        for row in sheet.iter_rows():
            for cell in row:
                if cell.value:
                    text = str(cell.value).strip()
                    extracted_text.append(text)

                    # 감지된 항목 추출
                    result = detect_sensitive_items_fullmatch(text)

                    for item, values in result.items():
                        detected_info.setdefault(item, set()).update(values)

    # set → list 변환
    detected_info = {k: list(v) for k, v in detected_info.items()}
    print(detected_info)
    return extracted_text, detected_info

# 셀 값에서 개인정보를 감지하고 마스킹된 값을 반환하는 함수
def mask_cell_value(value, selected_items):
    text = str(value)
    detected = detect_sensitive_items_fullmatch(text)

    for item_type, item_texts in detected.items():
        for detected_text in item_texts:
            if detected_text not in selected_items:
                continue

            # 감지된 텍스트가 전체 셀 텍스트와 같으면 전체 마스킹
            if detected_text == text:
                if item_type == "이름":
                    return '*' * len(text)
                elif item_type == "전화번호":
                    return ''.join(c if c == '-' else '*' for c in text)
                elif item_type == "이메일":
                    return ''.join(c if c in {'@', '.'} else '*' for c in text)
                elif item_type == "주민등록번호":
                    return ''.join(c if c == '-' else '*' for c in text)
                elif item_type == "주소":
                    return ''.join(c if c in {',', '-'} else '*' for c in text)
                else:
                    return '*' * len(text)  # 기타 항목

            # 셀에 여러 정보가 포함된 경우 해당 부분만 마스킹
            elif detected_text in text:
                masked = '*' * len(detected_text)
                text = text.replace(detected_text, masked)

    return text

# 엑셀 파일에서 개인정보 항목을 탐지하고 마스킹하여 저장하는 함수
def mask_excel_file(input_path, output_path, selected_items):
    import openpyxl
    import os

    workbook = openpyxl.load_workbook(input_path)

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]

        for row in sheet.iter_rows():
            for cell in row:
                if cell.value:
                    masked_value = mask_cell_value(cell.value, selected_items)
                    cell.value = masked_value

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    workbook.save(output_path)
    return output_path

